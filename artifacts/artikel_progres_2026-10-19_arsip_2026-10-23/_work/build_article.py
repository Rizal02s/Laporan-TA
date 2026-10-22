from pathlib import Path
import csv, json, zipfile, re
from collections import Counter
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.graphics.shapes import Drawing, String
from reportlab.graphics.charts.barcharts import HorizontalBarChart
from reportlab.graphics import renderPDF
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from pdf2image import convert_from_path

ROOT = Path(r'C:\VAEP-Track')
OUT = ROOT/'docs'/'artikel_progres_2026-10-19'
WORK = OUT/'_work'
OUT.mkdir(parents=True, exist_ok=True)
WORK.mkdir(exist_ok=True)
TITLE = 'Pengembangan Temporal Graph Neural Network untuk Deteksi Gegenpressing pada Data Spatiotemporal IDSSE'
AUTHOR = 'Rizal Andi Pranata'
AFFIL = 'Program Studi Teknik Informatika, Institut Asia Malang'
DATE = '19 Oktober 2026'

def readcsv(name):
    with (ROOT/'data'/'processed'/name).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))
ann = readcsv('gegenpressing_annotation_manifest.csv')
episodes = readcsv('episode_manifest.csv')
labels = Counter(r['is_gegenpressing'] for r in ann)
assert (len(ann), labels['1'], labels['0'], labels['-1'], labels['']) == (264,105,16,11,132)
assert len(episodes)==2353 and sum(r['strict_valid'].lower()=='true' for r in episodes)==1873
SPLITS = {'J03WMX':'Train','J03WN1':'Train','J03WOH':'Train','J03WOY':'Train','J03WR9':'Train','J03WQQ':'Validation','J03WPY':'Test'}

ABSTRACT = ('Deteksi gegenpressing memerlukan pengamatan respons kolektif pemain segera setelah kehilangan penguasaan bola. Penelitian ini mengembangkan pipeline deteksi berbasis Temporal Graph Neural Network pada Integrated Dataset of Spatiotemporal and Event Data in Elite Soccer (IDSSE) serta menyiapkan logistic regression dan XGBoost sebagai pembanding. Pendekatan yang digunakan adalah penelitian kuantitatif dengan eksperimen komputasional komparatif dan pembelajaran terawasi. Data tujuh pertandingan diolah menjadi kandidat transisi, disaring berdasarkan kelengkapan tracking, kemudian direpresentasikan sebagai urutan graph pemain dan bola. Setiap episode mencakup dua detik sebelum hingga dua detik sesudah kehilangan bola pada frekuensi 25 Hz, menghasilkan 101 frame dengan 23 node dan enam fitur per node. Arsitektur model menggabungkan dua lapisan Graph Convolutional Network, pooling per frame, dan Gated Recurrent Unit. Pemeriksaan hingga 19 Oktober 2026 menghasilkan 2.353 kandidat dan 1.873 episode valid ketat. Dari 264 episode dalam antrean anotasi, 132 telah diisi, terdiri atas 105 gegenpressing, 16 bukan gegenpressing, dan 11 ragu; 121 label biner tersedia. Pipeline telah menjalani uji integrasi pelatihan dan evaluasi, tetapi hasil tersebut berasal dari 26 label pilot. Anotasi validation dan test masing-masing baru mencakup tiga episode sehingga evaluasi akhir belum layak dilakukan. Capaian saat ini mendukung kelayakan teknis pembentukan dataset dan model. Keunggulan Temporal GNN masih memerlukan penyelesaian anotasi, audit label, penyelarasan representasi spasial, dan eksperimen akhir berbasis pemisahan pertandingan.')
KEYWORDS = 'gegenpressing; graph neural network; data tracking; IDSSE; anotasi manual'
ABSTRACT_EN = ('Gegenpressing detection requires observing collective player responses immediately after possession loss. This study develops a temporal graph neural network pipeline using the Integrated Dataset of Spatiotemporal and Event Data in Elite Soccer (IDSSE), with logistic regression and XGBoost as comparative models. The study adopts a quantitative, comparative computational design with supervised learning. Tracking data from seven matches are converted into transition candidates, screened for structural completeness, and represented as player–ball graph sequences. Each episode spans two seconds before and two seconds after possession loss at 25 Hz, yielding 101 frames, 23 nodes, and six node features. The architecture combines two graph convolutional layers, frame-level pooling, and a gated recurrent unit. As of 19 October 2026, the pipeline identified 2,353 candidates and 1,873 strictly valid episodes. Of 264 selected episodes, 132 have been annotated: 105 positive, 16 negative, and 11 uncertain, leaving 121 binary labels. Training and evaluation integration tests were completed using 26 pilot labels. However, the current validation and test sets each contain only three annotated episodes, preventing a meaningful final evaluation. These findings establish the technical feasibility of dataset preparation and model execution. Any claim of improved detection performance requires complete annotation, label auditing, spatial-representation checks, and final match-level evaluation.')

SECTIONS = {}
def P(t): return ('p',t)
def H(t): return ('h',t)
def T(caption,headers,rows,widths): return ('table',caption,headers,rows,widths)
def F(path,caption): return ('figure',str(path),caption)

SECTIONS[1] = [
P('Transisi setelah kehilangan penguasaan bola merupakan fase penting dalam analisis taktik sepak bola. Respons pemain pada fase ini dapat berupa tekanan segera terhadap penguasaan lawan atau penataan kembali organisasi bertahan. Gegenpressing, yang juga disebut counterpressing, berkaitan dengan respons tekanan setelah kehilangan bola. Identifikasinya memerlukan pembacaan posisi, arah pendekatan, keterlibatan pemain, dan keberlanjutan tekanan dalam rentang waktu singkat. Bauer dan Anzer [3] menempatkan identifikasi perilaku ini sebagai masalah klasifikasi terawasi berdasarkan data posisi dan event.'),
P('Perilaku gegenpressing perlu dipisahkan secara operasional dari hasil perebutan kembali penguasaan. Tekanan dapat dilakukan tanpa menghasilkan regain, sementara regain dapat terjadi karena kesalahan lawan. Apabila keberhasilan merebut bola langsung dipakai sebagai label perilaku, model berisiko mempelajari outcome yang berbeda dari tujuan deteksi. Penelitian ini karena itu menggunakan anotasi perilaku sebagai target, sedangkan quick regain dipertahankan sebagai variabel outcome untuk analisis lanjutan.'),
P('Data posisi berurutan menyediakan informasi yang relevan untuk menggambarkan hubungan pemain dengan bola dan lawan. IDSSE menyediakan data tersebut bersama catatan event dan informasi pertandingan dalam sumber yang dapat diakses publik [1], [2]. Penggunaannya memungkinkan proses pembentukan episode dan representasi model diperiksa ulang. Namun, akses publik dan kelengkapan koordinat tidak dengan sendirinya menyediakan label gegenpressing atau menjamin kemampuan generalisasi model.'),
P('Model berbasis fitur ringkas dapat merangkum jumlah penekan, jarak terhadap bola, dan durasi tekanan. Rangkuman tersebut memudahkan pembelajaran model tabular, tetapi tidak mempertahankan seluruh susunan relasional pada setiap frame. Penelitian ini menguji pendekatan yang menyajikan pemain dan bola sebagai graph berurutan. Graph Convolutional Network (GCN) digunakan untuk mengolah hubungan lokal, sedangkan Gated Recurrent Unit (GRU) digunakan untuk mempelajari perubahan representasi sepanjang episode [4], [5].'),
P('Pertanyaan penelitian adalah bagaimana membentuk episode transisi IDSSE menjadi graph temporal yang konsisten, bagaimana menyediakan label perilaku yang dapat diaudit, dan apakah pendekatan GCN–GRU memberikan hasil deteksi yang berbeda dari logistic regression serta XGBoost pada pembagian pertandingan yang sama. Hipotesis kerja menyatakan bahwa informasi hubungan spasial dan urutan waktu berpotensi membantu deteksi. Hipotesis tersebut belum dinyatakan terbukti pada tahap penelitian ini.'),
P('Tujuan penelitian adalah mengembangkan pipeline pengolahan data dan anotasi, mengimplementasikan model deteksi beserta pembanding, dan menyiapkan evaluasi yang dapat direproduksi. Kontribusi yang telah tersedia mencakup manifest episode, representasi graph, perangkat anotasi, dan implementasi pelatihan serta evaluasi. Artikel ini melaporkan hasil pengolahan dan kesiapan eksperimen hingga 19 Oktober 2026. Ruang lingkup kesimpulan dibatasi pada capaian teknis dan kualitas data yang telah diperiksa; analisis efektivitas dan valuasi aksi ditempatkan sebagai pengembangan lanjutan.')]

SECTIONS[2] = [
H('Deteksi counterpressing berbasis data'),
P('Bauer dan Anzer [3] mengembangkan deteksi counterpressing menggunakan data posisi dan event, label manual, serta 134 fitur berbasis pengetahuan pakar. Penelitian tersebut menunjukkan peran pembelajaran terawasi dan model XGBoost dalam identifikasi transisi. Baseline pada penelitian ini menggunakan 12 fitur tekanan yang tersedia dalam implementasi lokal. Dengan demikian, baseline tersebut merupakan pembanding yang disesuaikan dengan ruang lingkup data, bukan replikasi seluruh fitur dan prosedur studi terdahulu. Perbedaan dataset, pedoman anotasi, dan pembagian data juga membuat angka performa antarpublikasi tidak dapat dibandingkan langsung.'),
H('Dataset posisi dan representasi graph'),
P('Bassek et al. [1] memperkenalkan IDSSE yang mencakup dua pertandingan Bundesliga dan lima pertandingan divisi kedua musim 2022/2023. Berkas aslinya terdiri atas informasi pertandingan, event, dan posisi dalam format XML. Data posisi dikumpulkan melalui sistem optik TRACAB Gen5 pada 25 Hz. Dataset tersedia dengan lisensi CC BY 4.0 [2], sedangkan video pertandingan tidak disertakan. Cakupan tujuh pertandingan mendukung pengembangan metode awal, tetapi membatasi interpretasi hasil untuk populasi pertandingan yang lebih luas.'),
P('GCN menggabungkan fitur simpul dengan informasi struktur graph melalui agregasi dari lingkungan lokal [4]. Untuk kasus ini, simpul mewakili pemain dan bola, sedangkan edge dibentuk menurut kedekatan posisi. Edge tersebut merupakan keputusan representasi yang disusun peneliti; keberadaannya tidak menyatakan bahwa umpan, tekel, atau tekanan benar-benar terjadi. Makna terhadap target gegenpressing dipelajari melalui hubungan antara representasi episode dan label acuan.'),
H('Pemodelan urutan dan model pembanding'),
P('GRU menggunakan mekanisme gerbang untuk mengatur pembaruan keadaan tersembunyi pada data berurutan [5]. Pada arsitektur yang dikembangkan, keluaran GCN diringkas menjadi vektor per frame sebelum diproses GRU. Istilah temporal merujuk pada pemanfaatan urutan frame tersebut, sehingga model mempertimbangkan perubahan situasi selama episode, bukan hanya konfigurasi satu frame.'),
P('Logistic regression digunakan sebagai pembanding linear pada log-odds probabilitas kelas, dengan implementasi scikit-learn [7]. XGBoost merupakan sistem gradient boosting berbasis pohon yang memodelkan hubungan nonlinier melalui penambahan pohon secara bertahap [6]. Kedua model menerima fitur tekanan yang sama, sementara GNN menerima graph berurutan. Perbandingan utama karena itu menilai pendekatan representasi beserta modelnya; eksperimen ablasi diperlukan untuk mengisolasi manfaat graph atau komponen temporal secara terpisah.'),
H('Deteksi perilaku dan analisis efektivitas'),
P('VAEP atau Valuing Actions by Estimating Probabilities menilai aksi melalui perubahan probabilitas mencetak dan kebobolan gol dengan mempertimbangkan konteks keadaan permainan [8]. Kerangka ini menjadi arah pengembangan proyek VAEP-Track setelah deteksi perilaku tersedia. Pada tahap sekarang, pemetaan event IDSSE ke skema aksi dan validasi model probabilitas VAEP belum diselesaikan. Oleh karena itu, hasil artikel ini belum mencakup skor VAEP atau kesimpulan mengenai efektivitas gegenpressing.'),
T('Tabel 1 Posisi penelitian terhadap landasan metode', ['Rujukan','Pokok pendekatan','Hubungan dengan penelitian ini'], [
['Bauer dan Anzer [3]','Deteksi counterpressing berlabel manual dan fitur pakar','Dasar target perilaku dan pembanding berbasis fitur'],
['Kipf dan Welling [4] serta Cho et al. [5]','Agregasi lokal graph dan pemodelan urutan','Landasan komponen GCN dan GRU'],
['Chen dan Guestrin [6]','Tree boosting','Landasan pembanding XGBoost'],
['Decroos et al. [8]','Valuasi aksi melalui probabilitas gol','Arah analisis lanjutan yang belum dievaluasi']],[3.1,5.4,7.5])]

SECTIONS[3] = [
H('Rancangan penelitian dan sumber data'),
P('Penelitian menggunakan pendekatan kuantitatif dengan eksperimen komputasional komparatif pada data pertandingan sekunder. Tugas utama adalah klasifikasi biner episode menjadi gegenpressing atau bukan gegenpressing melalui supervised learning. Unit analisis adalah episode kehilangan bola; frame dan pemain merupakan komponen dalam episode, bukan sampel penelitian yang berdiri sendiri. Hasil yang tersedia pada tahap ini berupa audit data, statistik anotasi, serta pengujian integrasi implementasi.'),
P('Data utama berasal dari IDSSE [1], [2]. Skrip proyek mengakses event dan tracking menggunakan library Kloppy, mengubah keduanya menjadi tabel, lalu menyimpannya sebagai CSV. Versi Kloppy yang diperiksa mengarahkan akses dataset terbuka ke mirror PySport yang merujuk repositori asli Figshare. Pada CSV tracking, satu baris menyatakan satu frame dan kolom menyimpan waktu, status bola, tim pemilik bola, serta koordinat dan kecepatan objek. Identitas pemain dipetakan ke tim melalui catatan event. Pemetaan yang tidak lengkap akan memengaruhi kelolosan audit roster.'),
H('Deteksi kehilangan bola dan validasi episode'),
P('Kandidat turnover ditentukan ketika identitas tim pemilik bola berubah dalam periode pertandingan yang sama, kedua identitas tersedia, dan kondisi bola memenuhi syarat hidup. Setiap kandidat memiliki identitas unik yang menghubungkan waktu kehilangan bola, pertandingan, tim kehilangan bola, serta tim lawan. Quick regain dalam lima detik dicatat secara terpisah dari label perilaku dan tidak dimasukkan sebagai fitur deteksi.'),
P('Window episode ditetapkan dari −2,00 sampai +2,00 detik relatif terhadap kehilangan bola. Frekuensi 25 Hz menghasilkan 100 interval atau 101 frame ketika kedua ujung disertakan. Episode dinyatakan valid ketat apabila window lengkap, nomor frame kontinu, interval waktu 0,04 detik, posisi objek valid, dan terdapat 22 pemain aktif dengan pembagian 11 lawan 11. Kebijakan ini menghasilkan input dengan ukuran seragam, sekaligus membatasi cakupan pada situasi yang memenuhi struktur tersebut.'),
H('Pemilihan sampel dan anotasi perilaku'),
P('Antrean anotasi dibentuk dari episode valid dengan target paling banyak 40 episode per pertandingan. Sampling menggunakan random seed 42, kuota seimbang menurut outcome quick regain, dan jarak waktu minimal dua detik pada periode yang sama. Keterbatasan kandidat yang memenuhi kuota menghasilkan 24 episode untuk J03WN1 dan 40 episode untuk masing-masing pertandingan lain, sehingga totalnya 264. Keseimbangan outcome tidak menyatakan keseimbangan kelas gegenpressing; distribusi sampel ini juga tidak dapat langsung digunakan untuk mengestimasi prevalensi alami perilaku.'),
P('Anotasi dilakukan oleh peneliti menggunakan visualisasi lima cuplikan posisi pada 0; 0,5; 1; 1,5; dan 2 detik setelah turnover serta ringkasan indikator tekanan. Label 1 digunakan untuk gegenpressing, 0 untuk bukan gegenpressing, dan −1 untuk ragu. Setiap keputusan disertai confidence, identitas annotator, dan catatan ketika diperlukan. Label ragu dikeluarkan dari klasifikasi biner hingga peninjauan menghasilkan keputusan yang dapat dipertanggungjawabkan. Anotasi kedua yang independen dan ukuran kesepakatan antarpenilai belum tersedia.'),
P('Indikator pendukung mencakup kedekatan terhadap bola, laju mendekat, jumlah penekan, kekompakan lokal, dan durasi tekanan. Konfigurasi saat ini menggunakan radius tekanan dekat 4,572 m, radius pendekatan 10 m, laju mendekat minimal 1 m/s, serta durasi berkelanjutan 0,4 detik. Angka tersebut membantu pembacaan situasi dan bukan aturan otomatis penetapan label. Karena indikator yang ditampilkan juga menjadi bagian fitur baseline, audit independen diperlukan untuk menilai kemungkinan keputusan annotator terlalu mengikuti indikator tersebut.'),
H('Representasi graph temporal'),
P('Pada setiap frame, graph Gt = (V, Et, Xt) dibentuk dengan 23 node yang mewakili 22 pemain dan bola. Identitas serta urutan node dipertahankan sepanjang episode. Matriks Xt memiliki enam fitur: koordinat x, koordinat y, kecepatan, indikator bola, indikator tim kehilangan bola, dan indikator tim lawan. Bentuk tensor satu episode adalah 101 × 23 × 6. Status indikator tim dihitung relatif terhadap tim yang kehilangan bola pada episode tersebut.'),
P('Edge dibentuk secara dinamis dengan menghubungkan setiap node ke empat tetangga terdekat, sehingga konstruksi awal memuat 92 edge terarah per frame. Implementasi menghitung jarak Euclidean pada koordinat ternormalisasi CSV. Karena itu, kedekatan pada graph saat ini tidak dinyatakan sebagai jarak meter. Sebaliknya, indikator anotasi dan baseline menggunakan konversi posisi berdasarkan asumsi lapangan 105 × 68 m. Konsistensi skala dan dimensi lapangan aktual perlu diperiksa sebelum eksperimen akhir karena dapat mengubah urutan tetangga. Model GCN saat ini menerima edge_index tanpa bobot jarak edge secara eksplisit.'),
H('Model deteksi dan pembanding'),
P('Temporal GNN menggunakan dua lapisan GCN dengan ukuran tersembunyi 32, aktivasi ReLU, dan layer normalization. Masked mean pooling merangkum fitur node menjadi vektor tiap frame. Urutan vektor tersebut diproses GRU berukuran tersembunyi 32, kemudian keadaan terakhir diteruskan melalui dropout 0,2 dan lapisan linear dua kelas. Softmax menghasilkan probabilitas gegenpressing. Normalisasi fitur kontinu x, y, dan kecepatan dipelajari hanya dari train, sedangkan indikator jenis node dan tim tetap biner.'),
P('Baseline menggunakan 12 fitur: maksimum jumlah penekan total, penekan langsung, dan pemain yang mendekat; durasi tekanan individual dan kolektif; indikator tekanan berkelanjutan individual dan kolektif; jarak minimum ke bola; laju pendekatan maksimum; maksimum pemain dalam radius 10 m; rerata kekompakan lokal; serta perubahan bentangan tim. Fitur ini dihitung pada respons 0 sampai +2 detik. Perbedaan dengan input GNN yang mencakup −2 sampai +2 detik merupakan bagian rancangan saat ini dan perlu diperhitungkan dalam interpretasi perbandingan.'),
T('Tabel 2 Konfigurasi yang tersedia untuk eksperimen akhir', ['Model','Representasi','Konfigurasi utama'],[
['Logistic regression','12 fitur tekanan','Imputasi median, standardisasi, class_weight balanced, maksimum 2.000 iterasi'],
['XGBoost','12 fitur tekanan','200 pohon, max_depth 3, learning_rate 0,05, subsample dan colsample_bytree 0,8, bobot sampel seimbang'],
['Temporal GNN','101 graph per episode','GCN 32 dan 32, GRU 32, dropout 0,2, AdamW, batch 8, maksimum 100 epoch, patience 15']],[3.0,3.6,9.4]),
P('Konfigurasi pada Tabel 2 merupakan pengaturan yang telah disiapkan dalam implementasi dan runner pelatihan akhir, bukan hasil optimasi yang telah membuktikan konfigurasi terbaik. Pelatihan GNN menggunakan cross-entropy berbobot kelas dan pemilihan checkpoint berdasarkan validation loss. Seluruh keputusan pengaturan dan perubahan konfigurasi perlu dicatat sebelum evaluasi akhir.'),
H('Pembagian data dan prosedur evaluasi'),
P('Pembagian dilakukan menurut pertandingan: J03WMX, J03WN1, J03WOH, J03WOY, dan J03WR9 sebagai train; J03WQQ sebagai validation; serta J03WPY sebagai test. Dengan cara ini, episode dari pertandingan yang sama tidak berada pada split berbeda. Kandidat, label, dan identitas split dibuat konsisten untuk ketiga model. Imputasi dan normalisasi dipelajari dari train. Ambang keputusan tiap model dipilih menggunakan F1 pada validation, lalu digunakan pada test tanpa pemilihan ulang berdasarkan hasil test.'),
P('Evaluasi direncanakan menggunakan precision, recall, F1, ROC-AUC, average precision (AP), serta confusion matrix. Precision adalah TP/(TP + FP), recall adalah TP/(TP + FN), dan F1 adalah 2 × precision × recall/(precision + recall). AP merangkum hubungan precision–recall pada berbagai ambang. Implementasi saat ini menyimpan AP dalam kolom bernama pr_auc; artikel menggunakan istilah AP agar tidak menyamakannya dengan integrasi trapezoidal kurva precision–recall. Hasil perlu ditafsirkan bersama jumlah sampel dan distribusi kelas.'),
P('Runner membatasi pelatihan final hingga seluruh antrean anotasi diisi, identitas kandidat dan label valid, serta setiap split memiliki sekurang-kurangnya lima contoh untuk masing-masing kelas. Batas lima merupakan pengaman operasional pipeline, bukan pembuktian kecukupan statistik. Snapshot anotasi, konfigurasi, prediksi, dan checksum direncanakan disimpan sebagai artefak reproduksibilitas. Pada progres saat ini, syarat pelatihan akhir belum terpenuhi.')]

match_rows=[]
for m in sorted(SPLITS):
    ee=[r for r in episodes if r['match_id']==m]
    aa=[r for r in ann if r['match_id']==m]
    c=Counter(r['is_gegenpressing'] for r in aa)
    match_rows.append([m,SPLITS[m],str(len(ee)),str(sum(r['strict_valid'].lower()=='true' for r in ee)),str(len(aa)),str(len(aa)-c[''])])
match_rows.append(['Total','', '2.353','1.873','264','132'])
split_rows=[]
for s in ['Train','Validation','Test']:
    aa=[r for r in ann if SPLITS[r['match_id']]==s]
    c=Counter(r['is_gegenpressing'] for r in aa)
    split_rows.append([s,str(len(aa)),str(c['1']),str(c['0']),str(c['-1']),str(c['']),str(c['1']+c['0'])])
split_rows.append(['Total','264','105','16','11','132','121'])

pdfmetrics.registerFont(TTFont('Arial',r'C:\Windows\Fonts\arial.ttf'))
chart=HorizontalBarChart()
chart.x=135;chart.y=35;chart.width=330;chart.height=150
chart.data=[[132,11,16,105]]
chart.categoryAxis.categoryNames=['Belum dianotasi','Ragu','Bukan gegenpressing','Gegenpressing']
chart.categoryAxis.labels.fontName='Arial';chart.categoryAxis.labels.fontSize=10
chart.categoryAxis.labels.boxAnchor='e';chart.categoryAxis.labels.dx=-6
chart.valueAxis.valueMin=0;chart.valueAxis.valueMax=150;chart.valueAxis.valueStep=30
chart.valueAxis.labels.fontName='Arial';chart.valueAxis.labels.fontSize=9
chart.valueAxis.visibleGrid=True;chart.valueAxis.gridStrokeColor=colors.HexColor('#dddddd')
chart.bars[0].fillColor=colors.HexColor('#595959');chart.bars[0].strokeColor=None
chart.barLabelFormat='%d';chart.barLabels.fontName='Arial';chart.barLabels.fontSize=10
chart.barLabels.dx=9;chart.barLabels.boxAnchor='w'
drawing=Drawing(500,205);drawing.add(chart)
drawing.add(String(300,8,'Jumlah episode',fontName='Arial',fontSize=10,textAnchor='middle'))
renderPDF.drawToFile(drawing,str(WORK/'anotasi_chart.pdf'))
convert_from_path(str(WORK/'anotasi_chart.pdf'),dpi=200,poppler_path=r'C:\Users\rizal\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin')[0].save(WORK/'anotasi.png')

SECTIONS[4] = [
H('Ketersediaan dan kelolosan data'),
P('Pemeriksaan manifest terkini menghasilkan 2.353 kandidat defensive transition saat bola hidup. Sebanyak 1.873 episode atau 79,60% memenuhi kriteria valid ketat, sementara 480 episode tidak lolos. Tabel 3 menunjukkan bahwa tingkat kelolosan tidak merata antarlaga. J03WMX memiliki 408 dari 408 episode valid, sedangkan J03WN1 hanya 33 dari 394 atau 8,38%. J03WQQ memiliki 195 dari 304 atau 64,14% episode valid. Hasil ini menunjukkan perlunya melaporkan penyaringan per pertandingan karena jumlah kandidat awal belum menyatakan jumlah input yang dapat digunakan.'),
T('Tabel 3 Kandidat dan progres anotasi per pertandingan', ['Pertandingan','Split','Kandidat','Valid ketat','Antrean','Diisi'],match_rows,[2.5,2.8,2.7,2.8,2.6,2.6]),
P('Pemeriksaan valid ketat menguji roster 22 pemain, pembagian 11 lawan 11, kontinuitas frame, kelengkapan window, dan validitas koordinat. Episode yang tidak lolos tidak langsung dapat disebut data sumber rusak; kegagalan juga dapat berasal dari syarat representasi atau pemetaan pemain ke tim. Konsekuensinya, hasil penelitian berlaku pada subset yang memenuhi kriteria tersebut. Penyaringan yang lebih banyak pada pertandingan tertentu dapat memengaruhi variasi situasi yang tersedia untuk model.'),
H('Progres anotasi terbaru'),
P('Dari antrean 264 episode, sebanyak 132 telah diisi dan 132 belum diisi. Komposisi hasil anotasi adalah 105 gegenpressing, 16 bukan gegenpressing, dan 11 ragu. Dengan pengecualian label ragu, tersedia 121 label biner atau 45,83% dari antrean. Persentase 50% menyatakan penyelesaian pengisian anotasi, bukan penyelesaian seluruh penelitian. Seluruh keputusan yang tersimpan saat ini berasal dari satu annotator.'),
F(WORK/'anotasi.png','Gambar 1 Distribusi status anotasi pada 264 episode hingga 19 Oktober 2026'),
T('Tabel 4 Distribusi anotasi menurut split pertandingan', ['Split','Antrean','Positif','Negatif','Ragu','Belum','Biner'],split_rows,[3.0,2.2,2.1,2.1,2.0,2.3,2.3]),
P('Sebagian besar anotasi selesai berada pada train. Validation dan test masing-masing baru memiliki dua label positif dan satu negatif. Keberadaan kedua kelas belum cukup untuk estimasi performa yang stabil. Penyelesaian antrean karena itu perlu mencakup seluruh pertandingan, dilanjutkan audit distribusi dan konsistensi label, bukan hanya penambahan label pada train.'),
P('Proporsi positif mencapai 86,78% dari 121 label biner saat ini. Sebagai ilustrasi, prediksi semua episode sebagai positif akan memperoleh akurasi sebesar proporsi tersebut pada himpunan anotasi ini, walaupun seluruh contoh negatif salah diklasifikasikan. Perhitungan ini bukan hasil model, melainkan penjelasan mengapa akurasi saja tidak memadai. Distribusi tersebut juga belum mewakili seluruh pertandingan karena anotasi belum selesai dan sampling dikondisikan pada outcome quick regain.'),
H('Kelayakan representasi dan integrasi model'),
P('Representasi standar memuat 101 frame, 23 node, dan enam fitur node, dengan 92 edge terarah pada konstruksi awal setiap frame. Graph cache, tabel fitur baseline, dan penggabungan kandidat dengan label telah disiapkan untuk antrean anotasi. Pengujian forward pass dan uji pelatihan menunjukkan bahwa data dapat melewati arsitektur serta menghasilkan keluaran prediksi. Capaian ini membuktikan keterhubungan komponen perangkat lunak pada sampel pengujian, belum membuktikan ketepatan deteksi secara ilmiah.'),
P('Artefak pelatihan yang tersedia merupakan smoke test menggunakan 26 label pilot: 20 train, tiga validation, dan tiga test. Artefak tersebut berbeda dari manifest anotasi terbaru yang berisi 121 label biner. Tidak terdapat artefak pelatihan final yang sudah selesai dan tervalidasi pada saat pemeriksaan. Oleh karena itu, artikel ini tidak menyajikan angka smoke test sebagai hasil perbandingan akhir, tidak menetapkan model terbaik, dan tidak menarik kesimpulan tentang keunggulan GNN.'),
P('Pemeriksaan kesiapan yang dijalankan pada manifest terbaru menghasilkan NOT_READY. Penyebabnya adalah antrean anotasi belum lengkap dan jumlah minimal contoh per kelas pada setiap split belum terpenuhi. Kondisi ini menempatkan penelitian pada tahap pengembangan dan pematangan data berlabel, dengan perangkat eksperimen telah tersedia tetapi evaluasi akhir masih menunggu penyelesaian prasyarat.'),
H('Pembahasan metodologis dan keterbatasan'),
P('Dukungan terhadap representasi graph bersifat konseptual dan teknis: data memiliki identitas objek, koordinat, serta urutan waktu yang dapat dibentuk menjadi hubungan lokal. Untuk menyatakan manfaat empiris, eksperimen masih harus menunjukkan apakah representasi tersebut membantu dibandingkan fitur ringkas. Perbedaan input juga perlu dipertahankan dalam pembahasan: baseline memakai indikator pada 0 sampai +2 detik, sedangkan GNN menggunakan konteks −2 sampai +2 detik. Penyeragaman window atau ablasi konteks diperlukan apabila tujuan analisis adalah mengisolasi sumber peningkatan performa.'),
P('Audit skala koordinat menjadi pekerjaan metodologis sebelum eksperimen akhir. Jarak pada graph dihitung dalam bidang koordinat ternormalisasi, sedangkan fitur tekanan mengonversi posisi ke asumsi ukuran lapangan dalam meter. Karena panjang dan lebar lapangan tidak sama, urutan tetangga dalam ruang ternormalisasi dapat berbeda dari urutan jarak fisik. Pemeriksaan dimensi lapangan, konsistensi satuan kecepatan, dan pengaruh pilihan ruang jarak perlu didokumentasikan. Hal ini merupakan keterbatasan implementasi saat ini, bukan hasil pengujian sensitivitas yang sudah selesai.'),
P('Reliabilitas label belum diukur antarpengamat. Tampilan indikator tekanan dapat membantu konsistensi sekaligus memengaruhi penilaian, terutama karena sebagian indikator identik dengan fitur pembanding. Audit lanjutan perlu meninjau sampel ambigu dan sampel positif secara independen, dengan outcome quick regain disembunyikan dari penilai. Label perilaku harus tetap didasarkan pada respons kolektif yang diamati. Keseimbangan outcome pada pemilihan sampel juga perlu dilaporkan agar proporsi label dan metrik tidak ditafsirkan sebagai gambaran prevalensi alami.'),
P('Tujuh pertandingan membatasi variasi tim, konteks, dan kondisi permainan yang tercakup. Pemisahan pertandingan mengurangi risiko pencampuran pola laga antarsplit, tetapi satu pertandingan test tidak cukup untuk menyimpulkan kestabilan lintas liga. Episode yang berdekatan dalam satu pertandingan juga dapat saling berkaitan; jarak sampling minimal dua detik tidak menghilangkan seluruh overlap window empat detik. Penambahan evaluasi lintas pertandingan atau analisis sensitivitas dapat memperkuat kesimpulan setelah anotasi tersedia.'),
P('Input yang memuat dua detik sesudah turnover berarti deteksi memerlukan observasi respons setelah kehilangan bola. Klaim prediksi tepat saat turnover atau sistem waktu nyata tanpa latensi belum didukung oleh rancangan ini. Analisis hubungan gegenpressing dengan regain kelak juga bersifat observasional dan tidak dengan sendirinya menunjukkan sebab-akibat. Pengembangan VAEP memerlukan validasi tambahan atas skema aksi dan probabilitas gol yang berada di luar hasil yang telah dilaporkan.'),
H('Tahap eksperimen berikutnya'),
P('Pekerjaan berikutnya adalah menyelesaikan 132 anotasi tersisa, meninjau 11 label ragu, serta mengaudit definisi dan ketimpangan kelas. Setelah itu, satuan spasial dan kesetaraan konteks input perlu diperiksa, manifest terbaru dibekukan, dan pelatihan ketiga model dijalankan dengan keputusan parameter yang terdokumentasi. Evaluasi akhir mencakup metrik, jumlah sampel, confusion matrix, dan analisis kesalahan. Pengujian ablasi GCN tanpa GRU serta evaluasi lintas pertandingan dapat dipertimbangkan sebagai penguatan; keduanya belum menjadi hasil pada artikel ini.')]

SECTIONS[5] = [
P('Penelitian telah menghasilkan pipeline pengolahan IDSSE untuk deteksi gegenpressing, mencakup identifikasi kehilangan bola, pembentukan episode, anotasi manual, representasi graph temporal, serta implementasi model dan evaluasi. Sebanyak 2.353 kandidat menghasilkan 1.873 episode valid ketat. Representasi standar 101 frame dengan 23 node dan enam fitur per node dapat dibentuk secara konsisten, sehingga kelayakan teknis penggunaan data tracking untuk model graph telah ditunjukkan pada lingkup pengujian yang dilakukan.'),
P('Progres anotasi mencapai 132 dari 264 episode, dengan 105 label gegenpressing, 16 bukan gegenpressing, dan 11 ragu. Ketersediaan 121 label biner belum memenuhi kesiapan eksperimen akhir karena validation dan test masing-masing masih berisi tiga episode berlabel. Uji integrasi yang telah dijalankan menggunakan 26 label pilot hanya mendukung keberfungsian pipeline. Belum tersedia dasar untuk menetapkan keunggulan Temporal GNN terhadap logistic regression atau XGBoost.'),
P('Kesimpulan saat ini terbatas pada kelayakan pengolahan data dan implementasi metode. Penyelesaian anotasi, validasi label, pemeriksaan skala koordinat, dan evaluasi akhir berbasis pertandingan menjadi prasyarat untuk menilai manfaat pendekatan secara empiris. Cakupan tujuh pertandingan dan anotasi oleh satu penilai perlu dipertahankan sebagai batas generalisasi. Analisis efektivitas serta valuasi VAEP memerlukan tahapan tambahan dan belum menjadi temuan penelitian ini.')]

REFS = [
'[1] M. Bassek, R. Rein, H. Weber, dan D. Memmert, “An integrated dataset of spatiotemporal and event data in elite soccer,” Scientific Data, vol. 12, artikel 195, 2025. https://doi.org/10.1038/s41597-025-04505-y',
'[2] M. Bassek, R. Rein, H. Weber, dan D. Memmert, “An integrated dataset of spatiotemporal and event data in elite soccer,” Figshare, dataset, versi 1, 2025. https://doi.org/10.6084/m9.figshare.28196177.v1',
'[3] P. Bauer dan G. Anzer, “Data-driven detection of counterpressing in professional football,” Data Mining and Knowledge Discovery, vol. 35, hlm. 2009–2049, 2021. https://doi.org/10.1007/s10618-021-00763-7',
'[4] T. N. Kipf dan M. Welling, “Semi-supervised classification with graph convolutional networks,” International Conference on Learning Representations, 2017. https://arxiv.org/abs/1609.02907',
'[5] K. Cho, B. van Merriënboer, C. Gulcehre, D. Bahdanau, F. Bougares, H. Schwenk, dan Y. Bengio, “Learning phrase representations using RNN encoder–decoder for statistical machine translation,” Proceedings of the 2014 Conference on Empirical Methods in Natural Language Processing, hlm. 1724–1734, 2014. https://doi.org/10.3115/v1/D14-1179',
'[6] T. Chen dan C. Guestrin, “XGBoost: A scalable tree boosting system,” Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, hlm. 785–794, 2016. https://doi.org/10.1145/2939672.2939785',
'[7] F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” Journal of Machine Learning Research, vol. 12, hlm. 2825–2830, 2011. https://www.jmlr.org/papers/v12/pedregosa11a.html',
'[8] T. Decroos, L. Bransen, J. Van Haaren, dan J. Davis, “Actions speak louder than goals: Valuing player actions in soccer,” Proceedings of the 25th ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, hlm. 1851–1861, 2019. https://doi.org/10.1145/3292500.3330758'
]
NAMES={1:'Pendahuluan',2:'Tinjauan Pustaka',3:'Metode',4:'Hasil dan Pembahasan',5:'Simpulan',6:'Daftar Pustaka dan Abstraksi'}
ROMANS={1:'I',2:'II',3:'III',4:'IV',5:'V',6:'VI'}

def font(style,size=11,bold=False):
    style.font.name='Arial'; style.font.size=Pt(size); style.font.bold=bold; style.font.color.rgb=RGBColor(0,0,0)
    pr=style.element.get_or_add_rPr(); rf=pr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts');pr.append(rf)
    for k in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+k),'Arial')
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']: rf.attrib.pop(qn('w:'+k),None)

def newdoc(part=None):
    d=Document(); sec=d.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
    sec.header_distance=Cm(1);sec.footer_distance=Cm(1)
    for s in d.styles:
        if s.type in (1,2):
            try: font(s)
            except Exception: pass
    norm=d.styles['Normal'];norm.paragraph_format.line_spacing=1.12;norm.paragraph_format.space_after=Pt(7)
    norm.paragraph_format.widow_control=True
    for name,size in [('Title',15),('Heading 1',12),('Heading 2',11)]:
        font(d.styles[name],size,True)
        pf=d.styles[name].paragraph_format;pf.space_before=Pt(12);pf.space_after=Pt(6);pf.keep_with_next=True
    font(d.styles['Caption'],10)
    d.styles['Caption'].paragraph_format.space_before=Pt(4)
    d.styles['Caption'].paragraph_format.space_after=Pt(5)
    font(d.styles['Footer'],9)
    d.core_properties.title=TITLE
    d.core_properties.author=AUTHOR
    d.core_properties.subject='Progres penelitian hingga '+DATE
    p=d.add_paragraph(TITLE,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(AUTHOR);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2)
    p=d.add_paragraph(AFFIL);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
    if part:
        p=d.add_paragraph('Bagian '+ROMANS[part]+' '+NAMES[part]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(4)
        for r in p.runs:r.bold=True
    p=d.add_paragraph('Cakupan hasil hingga '+DATE);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(12)
    for r in p.runs:r.font.size=Pt(9)
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    return d

def paragraph(d,text,style=None):
    p=d.add_paragraph(text,style)
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY if not style else WD_ALIGN_PARAGRAPH.LEFT
    return p

def cellprops(cell,fill):
    pr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);pr.append(shd)
    margins=OxmlElement('w:tcMar')
    for side in ['top','bottom','start','end']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    borders=OxmlElement('w:tcBorders')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)

def table(d,caption,headers,rows,widths):
    p=paragraph(d,caption,'Caption');p.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    for i,row in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        tr=t.rows[i]._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if i==0: tr.append(OxmlElement('w:tblHeader'))
        for j,(cell,txt,w) in enumerate(zip(cells,row,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            cellprops(cell,'E9E9E9' if i==0 else 'FFFFFF')
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.05
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(txt);r.font.name='Arial';r.font.size=Pt(9.2);r.bold=(i==0)
    d.add_paragraph().paragraph_format.space_after=Pt(0)

def blocks(d,items):
    for b in items:
        if b[0]=='p': paragraph(d,b[1])
        elif b[0]=='h': d.add_paragraph(b[1],'Heading 2')
        elif b[0]=='table': table(d,*b[1:])
        elif b[0]=='figure':
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(b[1],width=Cm(15.8))
            paragraph(d,b[2],'Caption')

def abstract(d,english=True):
    d.add_paragraph('Abstrak','Heading 1');paragraph(d,ABSTRACT)
    p=paragraph(d,'Kata kunci: '+KEYWORDS);p.paragraph_format.space_after=Pt(10)
    if english:
        d.add_paragraph('Abstract','Heading 1');paragraph(d,ABSTRACT_EN)
        paragraph(d,'Keywords: gegenpressing; graph neural network; tracking data; IDSSE; manual annotation')

def references(d):
    d.add_paragraph('Daftar Pustaka','Heading 1')
    for r in REFS:
        p=paragraph(d,r);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent=Cm(.65);p.paragraph_format.first_line_indent=Cm(-.65)
        p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(7)
        for run in p.runs:run.font.size=Pt(10)

full=newdoc()
abstract(full)
for n in range(1,6):
    full.add_paragraph(ROMANS[n]+' '+NAMES[n],'Heading 1')
    blocks(full,SECTIONS[n])
references(full)
fullpath=OUT/'Artikel_Lengkap_Gegenpressing_Rizal_2026-10-19.docx'
full.save(fullpath)
files=[fullpath]
for n in range(1,7):
    d=newdoc(n)
    if n==6:
        abstract(d);references(d)
    else:
        d.add_paragraph(ROMANS[n]+' '+NAMES[n],'Heading 1')
        blocks(d,SECTIONS[n])
    path=OUT/('Bagian_'+ROMANS[n]+'_'+NAMES[n].replace(' ','_')+'.docx')
    d.save(path);files.append(path)

report={'author':AUTHOR,'affiliation':AFFIL,'date':DATE,'files':[str(f) for f in files],
        'counts':dict(labels),'matched_rows':match_rows,'split_rows':split_rows,
        'word_count_full':len(' '.join(p.text for p in full.paragraphs).split())}
(WORK/'build_manifest.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
for f in files: print(f.name)
print('Words',report['word_count_full'])
