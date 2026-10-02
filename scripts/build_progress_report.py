from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PROCESSED = ROOT / "data" / "processed"
REPORT_PATH = ROOT / "docs" / "Laporan_Progres_Kelayakan_VAEP_Track_2026-10-01.docx"
ASSET_DIR = Path.home() / "AppData" / "Local" / "Temp" / "vaep_track_report_assets"

DARK_BLUE = "17365D"
MID_BLUE = "2F75B5"
LIGHT_BLUE = "DDEBF7"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"
GREEN = "2E7D32"
AMBER = "A65D03"
RED = "A61B1B"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_column_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def style_table(table, widths_cm: list[float] | None = None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths_cm:
        set_column_widths(table, widths_cm)
    set_repeat_table_header(table.rows[0])
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, DARK_BLUE)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
                    run.font.size = Pt(8.5)
                    if row_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(doc, headers: list[str], rows: list[list[str]], widths_cm=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = str(value)
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
    style_table(table, widths_cm)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)


def add_status_callout(doc, title: str, text: str, color: str = GREEN) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EDF6EE" if color == GREEN else "FFF4E5")
    set_cell_borders(cell, color=color, size="8")
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(color)
    paragraph.add_run("\n" + text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def main() -> None:
    episodes = pd.read_csv(PROCESSED / "episode_manifest.csv")
    annotations = pd.read_csv(PROCESSED / "gegenpressing_annotation_manifest.csv")
    split_manifest = pd.read_csv(PROCESSED / "match_split_manifest.csv")

    match_summary = (
        episodes.groupby("match_id")
        .agg(
            candidates=("candidate_id", "size"),
            quick_regain=("label_regain_5s", "sum"),
            strict_valid=("strict_valid", "sum"),
        )
        .reset_index()
    )
    match_summary["strict_valid_rate"] = (
        100 * match_summary["strict_valid"] / match_summary["candidates"]
    ).round(1)
    pilot = annotations[
        annotations["pilot_selected"].astype(str).str.lower().eq("true")
        & annotations["is_gegenpressing"].notna()
    ].copy()
    pilot["is_gegenpressing"] = pilot["is_gegenpressing"].astype(int)
    quantity_figure = ASSET_DIR / "episode_quantity.png"
    annotation_figure = ASSET_DIR / "pilot_annotation.png"
    if not quantity_figure.exists() or not annotation_figure.exists():
        raise FileNotFoundError(
            "Grafik laporan belum dibuat. Jalankan build_progress_report_figures.py."
        )

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in (("Title", 21), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Heading 2" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    title_p_pr = styles["Title"].element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    header = section.header.paragraphs[0]
    header.text = "VAEP Track  |  Laporan Progres Proof of Concept"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("666666")
    add_page_number(section.footer.paragraphs[0])
    for run in section.footer.paragraphs[0].runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("666666")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Laporan Progres Proof of Concept VAEP Track")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Jawaban atas arahan kelayakan dataset IDSSE dan pipeline Graph Neural Network")
    run.bold = True
    run.font.size = Pt(12)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.add_run("1 Oktober 2026").italic = True

    add_status_callout(
        doc,
        "Kesimpulan utama",
        "Dataset IDSSE telah tersedia dan dapat dibaca, kuantitas transisi telah diaudit, "
        "graph temporal telah dibentuk secara konsisten, dan satu sequence telah berhasil "
        "melewati forward pass GNN. Topik layak dilanjutkan sebagai proof of concept dan "
        "evaluasi eksploratif, dengan anotasi manual penuh sebagai syarat sebelum training final.",
    )

    doc.add_heading("1 Tujuan laporan", level=1)
    doc.add_paragraph(
        "Laporan ini merangkum bukti teknis untuk menjawab empat arahan dosen pembimbing: "
        "ketersediaan data IDSSE, jumlah turnover yang dapat digunakan, konsistensi representasi "
        "graph, dan keberhasilan proof of concept GNN. Angka yang dilaporkan berasal dari output "
        "notebook dan manifest lokal yang telah dijalankan sampai selesai."
    )

    status_rows = [
        ["1", "Event dan tracking tersedia serta terbaca", "Terpenuhi", "7 pasangan file pertandingan; 1.002.644 frame tracking dan 10.964 baris event lokal."],
        ["2", "Jumlah turnover untuk training dan testing", "Terpenuhi dengan catatan", "2.353 kandidat aktif; 1.873 episode valid ketat. Label manual final belum lengkap."],
        ["3", "Tracking dapat diubah menjadi graph konsisten", "Terpenuhi", "23 node, 6 fitur node, 92 edge per frame, 101 frame per episode standar."],
        ["4", "Minimal satu sampel masuk pipeline GNN", "Terpenuhi", "Sequence PyTorch Geometric berhasil diproses; logit 2 kelas dan embedding 32 dimensi dihasilkan."],
    ]
    add_table(doc, ["Poin", "Arahan dosen", "Status", "Bukti utama"], status_rows, [1.0, 5.0, 3.0, 7.0])

    doc.add_heading("2 Ketersediaan dan keterbacaan dataset", level=1)
    doc.add_paragraph(
        "Dataset kerja memuat event dan tracking untuk tujuh pertandingan dengan ID J03WMX, "
        "J03WN1, J03WOH, J03WOY, J03WPY, J03WQQ, dan J03WR9. Seluruh frame memiliki informasi "
        "ball_owning_team_id yang dapat digunakan pada audit perubahan penguasaan. Pada audit "
        "J03WPY, tracking memiliki 146.211 baris dan 133 kolom, sedangkan event memiliki 1.572 "
        "baris dan 19 kolom."
    )
    add_bullet(doc, "Tracking menjadi sumber utama deteksi perubahan penguasaan karena tersedia kontinu pada 25 Hz.")
    add_bullet(doc, "Event digunakan sebagai konteks kejadian, bukan satu-satunya sumber turnover, karena event RECOVERY jarang dan sebagian timestamp periode kedua memerlukan penanganan khusus.")
    add_bullet(doc, "Data lokal dapat dibaca ulang tanpa error melalui pipeline Python proyek.")

    doc.add_heading("3 Kuantitas sampel transisi", level=1)
    doc.add_paragraph(
        "Deteksi menghasilkan 2.847 perubahan penguasaan secara keseluruhan. Setelah membatasi "
        "pada bola hidup, terdapat 2.353 kandidat defensive transition. Sebanyak 1.105 kandidat "
        "atau 46,96 persen diikuti quick regain dalam lima detik. Audit struktur menghasilkan "
        "1.873 episode valid ketat dan 480 episode tidak valid."
    )
    quantity_rows = []
    for row in match_summary.sort_values("match_id").itertuples(index=False):
        quantity_rows.append([
            row.match_id,
            f"{row.candidates:,}".replace(",", "."),
            f"{row.quick_regain:,}".replace(",", "."),
            f"{row.strict_valid:,}".replace(",", "."),
            f"{row.strict_valid_rate:.1f}%".replace(".", ","),
        ])
    quantity_rows.append(["Total", "2.353", "1.105", "1.873", "79,6%"])
    add_table(doc, ["Pertandingan", "Kandidat", "Quick regain", "Valid ketat", "Tingkat valid"], quantity_rows, [3.2, 2.8, 3.0, 3.0, 3.2])

    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(quantity_figure), width=Inches(6.3))
    caption = doc.add_paragraph("Gambar 1. Kandidat transisi dan episode valid ketat per pertandingan")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(9)

    add_status_callout(
        doc,
        "Catatan risiko kuantitas",
        "Jumlah episode valid secara struktural cukup untuk eksperimen, tetapi jumlah pertandingan "
        "tetap hanya tujuh. J03WN1 memiliki tingkat valid 8,4 persen dan J03WQQ 64,1 persen, "
        "terutama karena jumlah pemain aktif tidak selalu 22. Klaim penelitian harus dibatasi pada "
        "proof of concept dan evaluasi eksploratif.",
        color=AMBER,
    )

    doc.add_page_break()
    doc.add_heading("4 Konsistensi graph spatiotemporal", level=1)
    doc.add_paragraph(
        "Setiap episode standar menggunakan window empat detik, yaitu dua detik sebelum sampai dua "
        "detik setelah kehilangan bola. Dengan frekuensi 25 Hz, satu sequence berisi 101 graph. "
        "Representasi ini mempertahankan identitas dan susunan node pada seluruh frame."
    )
    graph_rows = [
        ["Node", "23", "11 pemain tim kehilangan bola, 11 lawan, dan 1 bola"],
        ["Fitur node", "6", "x, y, speed, is_ball, is_losing_team, is_opponent"],
        ["Edge", "92 per frame", "Empat tetangga spasial terdekat untuk setiap node"],
        ["Dimensi waktu", "101 frame", "Window -2,00 sampai +2,00 detik pada 25 Hz"],
        ["Validasi", "Ketat", "Window lengkap, frame kontinu, 25 Hz, pembagian pemain 11 lawan 11"],
    ]
    add_table(doc, ["Komponen", "Ukuran", "Definisi"], graph_rows, [3.2, 3.0, 9.8])

    doc.add_heading("5 Hasil proof of concept GNN", level=1)
    doc.add_paragraph(
        "Sampel J03WMX_P1_F12051 berhasil dibentuk menjadi sequence PyTorch Geometric sepanjang "
        "101 graph. Setiap graph memiliki x berukuran 23 x 6 dan edge_index berukuran 2 x 92. "
        "Sequence berhasil melewati forward pass pada perangkat CUDA dan menghasilkan logit dua "
        "kelas serta embedding 32 dimensi. Logit tersebut berasal dari model belum terlatih sehingga "
        "tidak diperlakukan sebagai prediksi atau ukuran akurasi."
    )
    poc_rows = [
        ["Candidate ID", "J03WMX_P1_F12051"],
        ["Panjang sequence", "101 graph"],
        ["Tensor fitur", "101 x 23 x 6"],
        ["Edge per frame", "92"],
        ["Output forward pass", "Logit 2 kelas; embedding 32 dimensi"],
        ["Status", "Berhasil sebagai bukti integrasi, belum training"],
    ]
    add_table(doc, ["Komponen", "Hasil"], poc_rows, [5.2, 10.8])

    doc.add_heading("6 Anotasi manual dan definisi target", level=1)
    doc.add_paragraph(
        "Pilot anotasi mencakup 28 transisi dari tujuh pertandingan dengan outcome quick regain "
        "seimbang, masing-masing 14 positif dan 14 negatif. Label perilaku gegenpressing tetap "
        "dipisahkan dari outcome quick regain. Hasil terkini terdiri atas 21 gegenpressing, lima "
        "bukan gegenpressing, dan dua ragu. Hanya 26 label biner yang dapat digunakan; dua label "
        "ragu dikeluarkan dari training."
    )
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(annotation_figure), width=Inches(5.4))
    caption = doc.add_paragraph("Gambar 2. Distribusi label manual pada pilot anotasi")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(9)
    doc.add_paragraph(
        "Pilot membuktikan bahwa prosedur anotasi dapat dijalankan dan disimpan, tetapi belum cukup "
        "untuk training atau estimasi generalisasi. Proporsi positif yang tinggi juga perlu diaudit "
        "pada anotasi lanjutan agar definisi tidak terlalu permisif."
    )

    doc.add_heading("7 Pembagian train validation dan test", level=1)
    doc.add_paragraph(
        "Pembagian dilakukan berdasarkan pertandingan, bukan berdasarkan frame atau episode, untuk "
        "mencegah potongan tracking dari pertandingan yang sama muncul pada lebih dari satu split. "
        "J03WMX ditempatkan di train karena telah digunakan selama pengembangan proof of concept."
    )
    split_rows = []
    split_names = {"train": "Train", "validation": "Validation", "test": "Test"}
    for split, group in split_manifest.groupby("split", sort=False):
        split_rows.append([
            split_names[split],
            ", ".join(group["match_id"].tolist()),
            f"{int(group['strict_valid_episodes'].sum()):,}".replace(",", "."),
            f"{100 * group['strict_valid_episodes'].sum() / split_manifest['strict_valid_episodes'].sum():.2f}%".replace(".", ","),
            str(int(group["annotated_usable"].sum())),
        ])
    add_table(doc, ["Split", "Pertandingan", "Episode valid", "Porsi", "Label pilot usable"], split_rows, [2.2, 5.6, 3.0, 2.2, 3.0])
    add_bullet(doc, "Train: 1.355 episode dari lima pertandingan.")
    add_bullet(doc, "Validation: 195 episode dari J03WQQ.")
    add_bullet(doc, "Test: 323 episode dari J03WPY.")
    add_bullet(doc, "Semua pemeriksaan candidate_id unik, satu split per pertandingan, label biner, dan eksklusi label ragu telah lulus.")

    doc.add_page_break()
    doc.add_heading("8 Keputusan kelayakan", level=1)
    doc.add_paragraph(
        "Keempat arahan awal dosen telah dijawab secara teknis. Dataset tersedia, jumlah kandidat "
        "telah dihitung, graph dapat dibentuk secara konsisten, dan satu sampel berhasil masuk ke "
        "pipeline GNN. Dengan demikian, topik dapat dilanjutkan sebagai Tugas Akhir dengan ruang "
        "lingkup proof of concept dan evaluasi eksploratif pada IDSSE."
    )
    doc.add_paragraph(
        "Kelayakan ini belum berarti model siap dilatih. Kendala terdekat adalah ketersediaan label "
        "manual. Model deteksi final, perbandingan baseline, dan valuasi efektivitas baru dilakukan "
        "setelah jumlah label memadai. VAEP kanonik juga memerlukan pemetaan event ke action schema; "
        "apabila pemetaan tidak dapat divalidasi, laporan akhir harus menggunakan istilah "
        "VAEP-inspired effectiveness score dan menjelaskan perbedaannya."
    )

    doc.add_heading("9 Rencana kerja berikutnya", level=1)
    next_steps = [
        "Mengunci pedoman anotasi dan menyelesaikan queue 264 transisi, dengan label ragu tetap dikeluarkan dari training.",
        "Melakukan audit kualitas label, termasuk peninjauan ulang sampel ambigu dan, bila memungkinkan, anotasi kedua pada subset untuk mengukur kesepakatan.",
        "Melatih baseline sederhana berbasis fitur sebelum temporal GNN agar manfaat representasi graph dapat diukur secara adil.",
        "Mengevaluasi model pada split berbasis pertandingan dengan precision, recall, F1, PR-AUC, ROC-AUC, dan confusion matrix.",
        "Menguji kompatibilitas event IDSSE dengan VAEP kanonik sebelum menetapkan metode valuasi akhir.",
    ]
    for step in next_steps:
        add_bullet(doc, step)

    doc.add_heading("10 Artefak verifikasi", level=1)
    artifacts = [
        ["01_data_audit.ipynb", "Audit event dan tracking serta keterbacaan data"],
        ["03_dataset_quantity.ipynb", "Jumlah perubahan penguasaan dan quick regain tujuh pertandingan"],
        ["04_gnn_poc_j03wmx.ipynb", "Forward pass GNN pada episode J03WMX"],
        ["05_episode_dataset_audit.ipynb", "Audit 2.353 kandidat dan 1.873 episode valid"],
        ["06_gegenpressing_annotation.ipynb", "Pilot anotasi manual 28 sampel"],
        ["07_research_readiness_and_split.ipynb", "Manifest split dan pemeriksaan bebas leakage"],
    ]
    add_table(doc, ["Artefak", "Fungsi"], artifacts, [6.8, 9.2])

    doc.add_heading("Referensi", level=1)
    references = [
        "Bassek, M., Rein, R., Weber, H., dan Memmert, D. (2025). An integrated dataset of spatiotemporal and event data in elite soccer. Scientific Data, 12, 195. https://doi.org/10.1038/s41597-025-04505-y",
        "Bauer, P., dan Anzer, G. (2021). Data-driven detection of counterpressing in professional football. Data Mining and Knowledge Discovery, 35, 2009-2049. https://doi.org/10.1007/s10618-021-00763-7",
        "Decroos, T., Bransen, L., Van Haaren, J., dan Davis, J. (2019). Actions Speak Louder Than Goals: Valuing Player Actions in Soccer. Proceedings of ACM SIGKDD. https://doi.org/10.1145/3292500.3330758",
    ]
    for reference in references:
        paragraph = doc.add_paragraph(reference)
        paragraph.paragraph_format.left_indent = Cm(0.7)
        paragraph.paragraph_format.first_line_indent = Cm(-0.7)
        paragraph.paragraph_format.space_after = Pt(4)

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    main()
