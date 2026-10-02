from __future__ import annotations

import hashlib
import shutil
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


REFERENCE = Path(
    r"C:\Users\rizal\OneDrive\Documents\Laporan PKL & TA\Laporan_Document_TA\Proposal_VAEP-Track.docx"
)
OUTPUT = Path(r"C:\VAEP-Track\docs\Proposal_Judul_VAEP-Track.docx")
REFERENCE_SHA256 = "2C1707FABBAF17F06D41C9205DB05E2F2C2BF36EDCAD9B9AC53117041A97FDC4"
EDITABLE_PARTS = {"word/document.xml", "docProps/core.xml"}

TITLE = (
    "VAEP-Track: Deteksi dan Analisis Efektivitas Gegenpressing pada Transisi "
    "Kehilangan Bola Menggunakan Temporal Graph Neural Network dan "
    "Spatiotemporal Tracking Data IDSSE"
)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def preserve_template_package(generated: Path) -> None:
    packaged = OUTPUT.with_suffix(".packaged.docx")
    with ZipFile(REFERENCE, "r") as source_zip, ZipFile(
        generated, "r"
    ) as generated_zip, ZipFile(packaged, "w", ZIP_DEFLATED) as output_zip:
        for item in source_zip.infolist():
            if item.filename in EDITABLE_PARTS:
                data = generated_zip.read(item.filename)
            else:
                data = source_zip.read(item.filename)
            output_zip.writestr(item, data)
    packaged.replace(OUTPUT)


def clear_inline_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_runs(paragraph, pieces: list[tuple[str, bool]]) -> None:
    clear_inline_content(paragraph)
    for text, bold in pieces:
        run = paragraph.add_run(text)
        run.bold = bold


def set_title(paragraph, text: str) -> None:
    clear_inline_content(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell(cell, text: str, *, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    clear_inline_content(paragraph)
    run = paragraph.add_run(text)
    run.bold = bold
    if bold:
        run.font.color.rgb = RGBColor(255, 255, 255)
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def fill_table(table, rows: list[list[str]]) -> None:
    if len(rows) != len(table.rows):
        raise ValueError(f"Expected {len(table.rows)} rows, received {len(rows)}")
    for row_index, (row, values) in enumerate(zip(table.rows, rows)):
        if len(values) != len(row.cells):
            raise ValueError("Table column count does not match the template")
        for cell, value in zip(row.cells, values):
            set_cell(cell, value, bold=row_index == 0)
        prevent_row_split(row)
    repeat_header(table.rows[0])


def main() -> None:
    actual_hash = file_sha256(REFERENCE)
    if actual_hash != REFERENCE_SHA256:
        raise RuntimeError(
            f"Reference changed: expected {REFERENCE_SHA256}, received {actual_hash}"
        )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    paragraphs = doc.paragraphs

    set_title(paragraphs[0], TITLE)
    set_runs(paragraphs[2], [("Judul yang direkomendasikan: ", True), (TITLE + ".", False)])

    set_runs(
        paragraphs[4],
        [
            ("Permasalahan: ", True),
            (
                "Gegenpressing adalah respons tekanan segera setelah sebuah tim kehilangan bola. "
                "Perilaku ini bersifat kolektif, berlangsung singkat, dan tidak identik dengan hasil "
                "akhir berupa penguasaan kembali. Tim dapat menekan dengan benar tetapi gagal merebut "
                "bola, sedangkan quick regain juga dapat terjadi karena kesalahan lawan tanpa tekanan "
                "kolektif. Karena itu penelitian harus memisahkan label perilaku gegenpressing dari "
                "outcome quick regain.",
                False,
            ),
        ],
    )
    set_runs(
        paragraphs[5],
        [
            ("Ide penelitian: ", True),
            (
                "VAEP-Track merepresentasikan pemain dan bola sebagai rangkaian graph selama empat "
                "detik di sekitar momen kehilangan penguasaan. Relasi spasial dipelajari oleh Graph "
                "Convolutional Network, sedangkan perubahan antarframe dirangkum oleh Gated Recurrent "
                "Unit. Target deteksi berasal dari anotasi manual perilaku gegenpressing. Quick regain, "
                "waktu menuju regain, dan nilai permainan dianalisis sesudah deteksi sebagai ukuran "
                "efektivitas yang terpisah.",
                False,
            ),
        ],
    )
    set_runs(
        paragraphs[6],
        [
            ("Alur konsep: ", True),
            (
                "Data tracking digunakan untuk mendeteksi perubahan ball owning team saat bola hidup. "
                "Setiap turnover dibentuk menjadi episode dari dua detik sebelum hingga dua detik "
                "sesudah kehilangan bola. Episode yang lolos audit struktur divisualisasikan untuk "
                "anotasi behavior, dibagi menurut pertandingan, lalu dipakai pada baseline logistic "
                "regression dan XGBoost serta temporal GNN. Prediksi gegenpressing kemudian dihubungkan "
                "dengan quick regain dan valuasi yang telah divalidasi.",
                False,
            ),
        ],
    )

    fill_table(
        doc.tables[0],
        [
            [
                "Event + Tracking",
                "Turnover + Episode",
                "Anotasi Behavior",
                "Baseline + Temporal GNN",
                "Outcome + Valuasi",
            ],
            [
                "Posisi, bola, identitas tim, dan konteks event",
                "Window -2 s hingga +2 s di sekitar kehilangan bola",
                "Gegenpressing, bukan, atau ragu",
                "Perbandingan fitur ringkas dengan sequence graph",
                "Quick regain, waktu regain, dan nilai permainan",
            ],
        ],
    )

    set_runs(
        paragraphs[8],
        [
            ("Dataset utama: ", True),
            (
                "IDSSE atau Integrated Dataset of Spatiotemporal and Event Data in Elite Soccer, "
                "dideskripsikan oleh Bassek et al. (2025) di Scientific Data. Dataset memuat event dan "
                "tracking dari tujuh pertandingan Bundesliga dan 2. Bundesliga musim 2022/23 pada "
                "frekuensi 25 Hz. Ketersediaan posisi pemain, bola, ball state, dan ball owning team "
                "memungkinkan turnover, episode transisi, serta graph pemain-bola dibentuk secara "
                "konsisten.",
                False,
            ),
        ],
    )
    set_runs(
        paragraphs[9],
        [
            ("Bukti kelayakan awal: ", True),
            (
                "Pipeline telah menemukan 2.353 kandidat defensive transition saat bola hidup. Sebanyak "
                "1.873 episode memenuhi syarat strict-valid dan 264 episode lintas tujuh pertandingan "
                "telah dipilih secara reproducible untuk anotasi behavior. Setiap sampel graph berisi "
                "101 frame, 23 node, enam fitur node, dan 92 edge terarah per frame. Hasil ini memenuhi "
                "permintaan proof of concept, tetapi belum menjadi bukti performa model final.",
                False,
            ),
        ],
    )

    fill_table(
        doc.tables[1],
        [
            ["Aspek", "Keterangan"],
            ["Sumber ilmiah", "Bassek et al. (2025), Scientific Data, DOI 10.1038/s41597-025-04505-y"],
            ["Kompetisi", "Bundesliga dan 2. Bundesliga musim 2022/23"],
            ["Cakupan dataset", "7 pertandingan, 207 pemain, 10 tim, 11.137 event, dan 1.002.644 frame posisi x/y"],
            ["Kandidat transisi", "2.353 perubahan penguasaan saat bola hidup; 1.105 quick regain dalam lima detik"],
            ["Episode strict-valid", "1.873 episode dengan window lengkap dan struktur graph konsisten"],
            ["Format graph", "101 frame x 23 node x 6 fitur; 92 edge terarah per frame"],
            ["Queue anotasi", "264 episode lintas tujuh pertandingan untuk label behavior"],
            ["Status saat disusun", "28 selesai: 21 gegenpressing, 5 bukan gegenpressing, 2 ragu; 236 tersisa"],
        ],
    )

    set_runs(
        paragraphs[11],
        [
            ("Bauer dan Anzer (2021) ", True),
            (
                "menjadi rujukan utama untuk deteksi counterpressing berbasis positional dan event data. "
                "Mereka menggunakan label manual, 134 fitur hasil rekayasa pengetahuan ahli, dan XGBoost "
                "untuk mengklasifikasikan defensive transition. Studi tersebut juga menunjukkan bahwa "
                "keberhasilan memperoleh kembali penguasaan perlu dianalisis sesudah perilaku tekanan "
                "diidentifikasi.",
                False,
            ),
        ],
    )
    set_runs(
        paragraphs[12],
        [
            ("Posisi VAEP-Track: ", True),
            (
                "penelitian ini mempertahankan unit analisis transisi dan kebutuhan ground truth manual, "
                "tetapi mengganti representasi fitur terancang dengan sequence graph yang mempertahankan "
                "relasi pemain-bola dan evolusi waktu. Logistic regression dan XGBoost tetap digunakan "
                "sebagai baseline. Kontribusi dinilai dari perbandingan yang adil pada kandidat dan split "
                "pertandingan yang sama, bukan dari kompleksitas model semata.",
                False,
            ),
        ],
    )

    fill_table(
        doc.tables[2],
        [
            ["Aspek", "Bauer dan Anzer (2021)", "VAEP-Track"],
            ["Data", "Position dan event data", "Event dan tracking IDSSE"],
            ["Unit analisis", "Defensive transition", "Episode turnover -2 s hingga +2 s"],
            ["Target deteksi", "Counterpressing berlabel manual", "Gegenpressing berlabel manual"],
            ["Representasi", "134 expert-engineered features", "Sequence graph 101 frame"],
            ["Model", "XGBoost dan pembanding", "Logistic regression, XGBoost, dan temporal GNN"],
            ["Outcome", "Recovery possession dalam lima detik", "Quick regain dan waktu menuju regain"],
            ["Valuasi", "Shots, xG, dan goals", "VAEP kanonik jika valid; jika tidak, skor VAEP-inspired"],
        ],
    )

    method_items = [
        (
            "Audit dan integrasi data: ",
            "memeriksa event, tracking, timestamp, frame, identitas pemain, posisi bola, ball state, dan ball owning team.",
        ),
        (
            "Deteksi turnover: ",
            "menandai perubahan pemilik bola hanya saat kedua ID tim tersedia dan bola berstatus hidup.",
        ),
        (
            "Pembentukan episode: ",
            "mengambil window -2,00 sampai +2,00 detik pada 25 Hz dan menyaring episode dengan 101 frame kontinu.",
        ),
        (
            "Anotasi behavior: ",
            "memberi label gegenpressing, bukan gegenpressing, atau ragu berdasarkan tekanan segera, approach rate, jumlah presser, compactness, dan durasi tekanan.",
        ),
        (
            "Graph construction: ",
            "membentuk 23 node untuk 22 pemain dan bola, enam fitur node, serta empat tetangga spasial terdekat per node pada setiap frame.",
        ),
        (
            "Split pertandingan: ",
            "memisahkan train, validation, dan test berdasarkan pertandingan agar frame atau pola dari laga yang sama tidak bocor antarsplit.",
        ),
        (
            "Baseline: ",
            "melatih logistic regression dan XGBoost dari 12 fitur tekanan yang dihitung hanya dari informasi episode.",
        ),
        (
            "Temporal GNN: ",
            "menggunakan GCN per frame, masked mean pooling, dan GRU untuk memprediksi probabilitas gegenpressing sepanjang sequence.",
        ),
        (
            "Evaluasi dan efektivitas: ",
            "memilih threshold pada validation, mengunci test, membandingkan precision, recall, F1, PR-AUC, ROC-AUC, lalu menganalisis quick regain dan menentukan kelayakan VAEP kanonik.",
        ),
    ]
    for paragraph, (lead, text) in zip(paragraphs[14:23], method_items):
        set_runs(paragraph, [(lead, True), (text, False)])

    set_runs(
        paragraphs[24],
        [
            (
                "IDSSE hanya mencakup tujuh pertandingan. Pembagian per pertandingan mengurangi leakage, "
                "tetapi tidak menghilangkan keterbatasan ukuran sampel. Kesimpulan akan dibatasi sebagai "
                "evaluasi metode dan proof of concept pada sampel IDSSE, bukan generalisasi untuk seluruh liga.",
                False,
            )
        ],
    )
    set_runs(
        paragraphs[25],
        [
            (
                "Training final menunggu anotasi 264 episode dan audit kualitas label selesai. Quick regain "
                "tidak digunakan sebagai pengganti label gegenpressing. Istilah VAEP kanonik hanya dipakai "
                "jika event IDSSE dapat dipetakan ke action schema dan model state-value tervalidasi; jika "
                "tidak, hasil dilaporkan sebagai VAEP-inspired effectiveness score dengan batasan eksplisit.",
                False,
            )
        ],
    )

    outputs = [
        "Dataset episode defensive transition yang terstandar, strict-valid, berlabel behavior, dan memiliki manifest split pertandingan.",
        "Dua baseline terverifikasi, yaitu logistic regression dan XGBoost, beserta prediksi dan metriknya.",
        "Model temporal GNN yang menerima sequence graph pemain-bola dan menghasilkan probabilitas gegenpressing.",
        "Evaluasi komparatif pada kandidat yang identik menggunakan precision, recall, F1, PR-AUC, ROC-AUC, confusion matrix, dan metrik per pertandingan.",
        "Analisis efektivitas melalui quick regain dan waktu regain, serta valuasi VAEP kanonik atau VAEP-inspired sesuai hasil validasi action schema.",
    ]
    for paragraph, text in zip(paragraphs[28:33], outputs):
        set_runs(paragraph, [(text, False)])

    properties = doc.core_properties
    properties.title = "Proposal Judul VAEP Track"
    properties.subject = "Deteksi dan analisis efektivitas gegenpressing menggunakan temporal GNN dan data IDSSE"
    properties.author = "Rizal"
    properties.keywords = "gegenpressing, IDSSE, tracking, temporal GNN, quick regain, VAEP"

    generated = OUTPUT.with_suffix(".generated.docx")
    doc.save(generated)
    preserve_template_package(generated)
    generated.unlink()
    print(OUTPUT)


if __name__ == "__main__":
    main()
